import csv
import io
import os
from abc import ABC, abstractmethod
from typing import Optional
from docx import Document
from pypdf import PdfReader


class BaseParser(ABC):
    """Abstract base class for document parsers."""

    @abstractmethod
    def parse(self, file_path: str) -> str:
        """Extract clean plain text from the file."""
        pass


class PDFParser(BaseParser):
    """Extracts text from digital PDFs with scanned document detection."""

    def parse(self, file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        reader = PdfReader(file_path)
        extracted_pages = []

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            stripped = page_text.strip()
            if stripped:
                extracted_pages.append(f"--- Page {page_idx + 1} ---\n{stripped}")

        full_text = "\n\n".join(extracted_pages).strip()

        # Scanned PDF Detection: if the PDF has pages but extracted text is virtually empty
        if len(reader.pages) > 0 and len(full_text) < 20:
            raise ValueError(
                "Scanned or image-only PDF detected. No extractable digital text found. OCR processing is required."
            )

        return full_text


class DocxParser(BaseParser):
    """Extracts text from modern DOCX Word files, including paragraphs and tables."""

    def parse(self, file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        doc = Document(file_path)
        content_parts = []

        # 1. Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # 2. Extract tables as Markdown-style tables
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append(" | ".join(row_cells))
            if table_lines:
                content_parts.append("\n".join(table_lines))

        full_text = "\n\n".join(content_parts).strip()
        if not full_text:
            raise ValueError("DOCX document is empty.")
        return full_text


class TextParser(BaseParser):
    """Extracts text from .txt, .md, and .csv with multi-encoding fallback."""

    ENCODINGS = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

    def parse(self, file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        raw_bytes = None
        with open(file_path, "rb") as f:
            raw_bytes = f.read()

        text = None
        for enc in self.ENCODINGS:
            try:
                text = raw_bytes.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if text is None:
            text = raw_bytes.decode("utf-8", errors="replace")

        # Normalize newlines for cross-platform consistency
        text = text.replace("\r\n", "\n")

        # If it's a CSV, format each row into clear comma-separated or readable text
        _, ext = os.path.splitext(file_path)
        if ext.lower() == ".csv":
            reader = csv.reader(io.StringIO(text))
            rows = [", ".join(row) for row in reader if any(field.strip() for field in row)]
            text = "\n".join(rows)

        clean_text = text.strip()
        if not clean_text:
            raise ValueError("Text document is empty.")
        return clean_text


def parse_document(file_path: str, mime_type: Optional[str] = None) -> str:
    """
    Dispatcher that selects the appropriate parser based on file extension and returns extracted text.
    """
    _, ext = os.path.splitext(file_path)
    ext_lower = ext.lower()

    if ext_lower == ".pdf":
        return PDFParser().parse(file_path)
    elif ext_lower == ".docx":
        return DocxParser().parse(file_path)
    elif ext_lower in {".txt", ".md", ".csv"}:
        return TextParser().parse(file_path)
    else:
        raise ValueError(f"Unsupported file format for text extraction: '{ext}'")
